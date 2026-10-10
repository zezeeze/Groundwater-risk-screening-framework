from __future__ import annotations

import json
import math
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook


PROJECT = Path(__file__).resolve().parents[2]
KB_XLSX = PROJECT / "data/rag/rag_knowledge_base.xlsx"
ENTERPRISE_KML = PROJECT / "data/rag/spatial_anonymized/anonymized_enterprise_boundaries_shifted.ovkml"
WELL_KML = PROJECT / "data/rag/spatial_anonymized/anonymized_wells_shifted.ovkml"
WELL_INPUT = PROJECT / "data/llm/structured_inputs_41_wells.jsonl"
CORPUS = PROJECT / "data/rag/rag_corpus.jsonl"
OUTPUT = PROJECT / "data/rag/well_enterprise_production_candidates.jsonl"

KML_NS = {"k": "http://www.opengis.net/kml/2.2"}
PROFILE_SOURCE_IDS = ["SRC_20", "SRC_19"]
SPATIAL_SOURCE_IDS = ["SRC_25", "SRC_24"]


def read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]


def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.write_text("\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + "\n", encoding="utf-8")


def extended_value(placemark: ET.Element, name: str) -> str | None:
    for item in placemark.findall(".//k:Data", KML_NS):
        if item.get("name") == name:
            node = item.find("k:value", KML_NS)
            return node.text.strip() if node is not None and node.text else None
    return None


def parse_coordinates(text: str) -> list[tuple[float, float]]:
    return [(float(parts[0]), float(parts[1])) for token in text.strip().split() if len(parts := token.split(",")) >= 2]


def load_geometries() -> tuple[dict[str, list[tuple[float, float]]], dict[str, tuple[float, float]]]:
    polygons: dict[str, list[tuple[float, float]]] = {}
    for pm in ET.parse(ENTERPRISE_KML).findall(".//k:Placemark", KML_NS):
        enterprise_id = extended_value(pm, "enterprise_id")
        coords = pm.find(".//k:Polygon//k:coordinates", KML_NS)
        if enterprise_id and coords is not None and coords.text:
            polygons[enterprise_id] = parse_coordinates(coords.text)

    wells: dict[str, tuple[float, float]] = {}
    for pm in ET.parse(WELL_KML).findall(".//k:Placemark", KML_NS):
        well_id = extended_value(pm, "well_id")
        coords = pm.find(".//k:Point/k:coordinates", KML_NS)
        if well_id and coords is not None and coords.text:
            wells[well_id] = parse_coordinates(coords.text)[0]
    return polygons, wells


def normalize_zone(value: object) -> str:
    text = str(value or "")
    return text[-1] if text.endswith(("A", "B")) else text


def split_items(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text or text in {"/", "无", "未列出", "None"}:
        return []
    return [x.strip() for x in re.split(r"[、,，;；]+", text) if x.strip() and x.strip() != "/"]


def normalized_indicator(value: str) -> str:
    text = re.sub(r"[\s（）()\[\]]", "", value).casefold()
    aliases = {
        "高锰酸盐指数": "耗氧量",
        "化学需氧量": "耗氧量",
        "codmn": "耗氧量",
        "cod": "耗氧量",
        "总溶解性固体": "溶解性总固体",
    }
    return aliases.get(text, text)


def load_knowledge_tables() -> tuple[dict[str, dict], dict[str, str], dict[str, str]]:
    workbook = load_workbook(KB_XLSX, read_only=True, data_only=True)
    profiles: dict[str, dict] = {}
    for row in workbook["匿名企业背景"].iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        profiles[str(row[0])] = {
            "enterprise_id": str(row[0]),
            "zone_id": normalize_zone(row[1]),
            "industry": row[2] or "",
            "products": row[3] or "",
            "raw_materials": row[4] or "",
            "operating_status": row[5] or "",
            "characteristic_pollutants": row[6] or "",
            "pollutant_data_status": row[7] or "",
            "priority_investigation_enterprise": row[8] or "",
            "screening_basis": row[9] or "",
            "profile_source_ids": PROFILE_SOURCE_IDS,
        }

    well_zones = {
        str(row[0]): normalize_zone(row[1])
        for row in workbook["研究井场地背景"].iter_rows(min_row=5, values_only=True)
        if row[0]
    }
    indicator_categories = {
        normalized_indicator(str(row[0])): str(row[3])
        for row in workbook["风险标签映射"].iter_rows(min_row=5, values_only=True)
        if row[0] and row[3]
    }
    return profiles, well_zones, indicator_categories


def projected(point: tuple[float, float], lat0: float) -> tuple[float, float]:
    lon, lat = point
    return lon * 111320.0 * math.cos(math.radians(lat0)), lat * 110540.0


def point_in_polygon(point: tuple[float, float], polygon: list[tuple[float, float]], lat0: float) -> bool:
    x, y = projected(point, lat0)
    vertices = [projected(item, lat0) for item in polygon]
    inside = False
    for (x1, y1), (x2, y2) in zip(vertices, vertices[1:] + vertices[:1]):
        if (y1 > y) != (y2 > y) and x < (x2 - x1) * (y - y1) / (y2 - y1) + x1:
            inside = not inside
    return inside


def point_segment_distance(point, start, end, lat0: float) -> float:
    x, y = projected(point, lat0)
    x1, y1 = projected(start, lat0)
    x2, y2 = projected(end, lat0)
    dx, dy = x2 - x1, y2 - y1
    denominator = dx * dx + dy * dy
    t = 0.0 if denominator == 0 else max(0.0, min(1.0, ((x - x1) * dx + (y - y1) * dy) / denominator))
    return math.hypot(x - (x1 + t * dx), y - (y1 + t * dy))


def boundary_distance(point, polygon, lat0: float) -> float:
    if point_in_polygon(point, polygon, lat0):
        return 0.0
    return min(point_segment_distance(point, a, b, lat0) for a, b in zip(polygon, polygon[1:] + polygon[:1]))


def polygon_centroid(polygon: list[tuple[float, float]]) -> tuple[float, float]:
    points = polygon[:-1] if polygon and polygon[0] == polygon[-1] else polygon
    return sum(x for x, _ in points) / len(points), sum(y for _, y in points) / len(points)


def relative_direction(well: tuple[float, float], enterprise: tuple[float, float]) -> str:
    dx = (enterprise[0] - well[0]) * math.cos(math.radians((enterprise[1] + well[1]) / 2))
    dy = enterprise[1] - well[1]
    angle = (math.degrees(math.atan2(dx, dy)) + 360) % 360
    labels = ["北", "东北", "东", "东南", "南", "西南", "西", "西北"]
    return labels[int((angle + 22.5) // 45) % 8]


def distance_band(distance: float) -> str:
    if distance == 0:
        return "企业边界内"
    if distance < 250:
        return "小于250 m"
    if distance < 500:
        return "250–500 m"
    if distance < 1000:
        return "500–1000 m"
    if distance < 2000:
        return "1–2 km"
    return "大于2 km"


def hydraulic_consistency(direction: str, internal: bool) -> str:
    if internal:
        return "企业内井，不适用上游方位筛选"
    if direction in {"北", "东北", "西北"}:
        return "区域流向支持潜在上游关系"
    if direction in {"东", "西"}:
        return "区域流向证据不足"
    return "区域流向不支持直接上游关系"


def category_probabilities(record: dict) -> dict[str, dict[str, float]]:
    return {
        str(item.get("category_name_zh") or item.get("category_id")): {
            "eem_probability": float(item.get("eem", {}).get("probability", 0.0)),
            "uvvis_probability": float(item.get("uvvis", {}).get("probability", 0.0)),
        }
        for item in record.get("risk_category_models", [])
    }


def production_alignment(profile: dict, category_map: dict[str, str], probabilities: dict[str, dict[str, float]]) -> dict:
    indicators = split_items(profile.get("characteristic_pollutants"))
    categories: list[str] = []
    matched_indicators: list[str] = []
    indicator_category_pairs: list[dict[str, str]] = []
    for indicator in indicators:
        normalized = normalized_indicator(indicator)
        category = category_map.get(normalized)
        if category is None:
            for key, candidate in category_map.items():
                if key and (key in normalized or normalized in key):
                    category = candidate
                    break
        if category:
            matched_indicators.append(indicator)
            indicator_category_pairs.append({"pollutant": indicator, "risk_category": category})
            if category not in categories:
                categories.append(category)
    details = []
    for category in categories:
        values = probabilities.get(category, {"eem_probability": 0.0, "uvvis_probability": 0.0})
        details.append({"category": category, **values})
    return {
        "reported_pollutants": indicators,
        "mapped_risk_categories": categories,
        "matched_reported_pollutants": matched_indicators,
        "reported_pollutant_category_map": indicator_category_pairs,
        "model_category_alignment": details,
        "alignment_status": "explicit characteristic-pollutant linkage available" if details else "no explicit characteristic-pollutant category linkage available",
    }


def relation_content(well_id: str, candidates: list[dict]) -> str:
    parts = []
    for item in candidates:
        location = "井位于企业边界内" if item["distance_band"] == "企业边界内" else f"距企业边界{item['distance_band']}，位于井的{item['enterprise_relative_direction']}侧"
        parts.append(
            f"{item['enterprise_id']}为第{item['distance_rank']}近候选，{location}，{item['regional_hydraulic_consistency']}；"
            f"行业={item['industry']}，生产状态={item['operating_status']}，特征污染指标={item['characteristic_pollutants'] or '未列出'}"
        )
    return "；".join(parts)


def main() -> None:
    profiles, well_zones, category_map = load_knowledge_tables()
    polygons, wells = load_geometries()
    records = read_jsonl(WELL_INPUT)
    by_well = {row["well_id"]: row for row in records}
    lat0 = sum(lat for _, lat in wells.values()) / len(wells)

    all_candidates: list[dict] = []
    for well_id, point in wells.items():
        distances = []
        for enterprise_id, polygon in polygons.items():
            profile = profiles.get(enterprise_id)
            if not profile or profile["zone_id"] != well_zones.get(well_id):
                continue
            distance = boundary_distance(point, polygon, lat0)
            internal = distance == 0
            direction = "井位于企业边界内" if internal else relative_direction(point, polygon_centroid(polygon))
            distances.append((distance, enterprise_id, direction, internal))
        distances.sort(key=lambda item: (item[0], item[1]))
        probabilities = category_probabilities(by_well[well_id])
        candidates = []
        for rank, (distance, enterprise_id, direction, internal) in enumerate(distances[:10], start=1):
            profile = profiles[enterprise_id]
            alignment = production_alignment(profile, category_map, probabilities)
            item = {
                "enterprise_id": enterprise_id,
                "distance_rank": rank,
                "distance_band": distance_band(distance),
                "enterprise_relative_direction": direction,
                "regional_hydraulic_consistency": hydraulic_consistency(direction, internal),
                "industry": profile["industry"],
                "operating_status": profile["operating_status"],
                "products": profile["products"],
                "raw_materials": profile["raw_materials"],
                "characteristic_pollutants": profile["characteristic_pollutants"],
                "pollutant_data_status": profile["pollutant_data_status"],
                "priority_investigation_enterprise": profile["priority_investigation_enterprise"],
                "production_alignment": alignment,
                "evidence_level": "spatial, hydraulic and production candidate evidence",
                "basis": "same-zone boundary distance and direction, regional groundwater flow, operating status, industrial activity, products, raw materials and reported characteristic pollutants",
                "limitations": "regional flow and enterprise activity constrain follow-up priority but do not establish contaminant origin",
                "source_ids": list(dict.fromkeys(SPATIAL_SOURCE_IDS + PROFILE_SOURCE_IDS)),
            }
            candidates.append(item)
            all_candidates.append({"well_id": well_id, **item})
        record = by_well[well_id]
        record["enterprise_context"] = {
            "candidate_pool_size": 10,
            "output_candidate_count": 5,
            "candidate_enterprises": candidates,
            "status": "spatially constrained candidates enriched with production evidence",
            "instruction": "select and rank five enterprises by jointly considering spatial and regional hydraulic relationships, operating status, industry, products, raw materials, reported characteristic pollutants and their consistency with the well-level risk-category evidence",
        }

    write_jsonl(WELL_INPUT, records)
    write_jsonl(OUTPUT, all_candidates)

    corpus = read_jsonl(CORPUS)
    relation_lookup = {well_id: by_well[well_id]["enterprise_context"]["candidate_enterprises"] for well_id in by_well}
    for entry in corpus:
        if entry.get("knowledge_type") == "井企候选关系" and entry.get("subject") in relation_lookup:
            candidates = relation_lookup[entry["subject"]]
            entry["content"] = relation_content(entry["subject"], candidates)
            entry["scope"] = "形成十家空间候选池，并为结合企业生产活动筛选五家优先核查企业提供证据"
            entry["limitations"] = "候选关系用于限定核查范围；企业活动、特征污染指标和区域水力关系均不能单独确定污染来源。"
            entry["source_ids"] = ";".join(list(dict.fromkeys(SPATIAL_SOURCE_IDS + PROFILE_SOURCE_IDS)))
            entry["updated"] = "2026-09-23"
    write_jsonl(CORPUS, corpus)

    print(json.dumps({
        "wells": len(records),
        "candidate_pool_per_well": 10,
        "candidate_records": len(all_candidates),
        "profiles_with_characteristic_pollutants": sum(bool(split_items(row["characteristic_pollutants"])) for row in all_candidates),
        "profiles_with_explicit_category_alignment": sum(bool(row["production_alignment"]["model_category_alignment"]) for row in all_candidates),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
